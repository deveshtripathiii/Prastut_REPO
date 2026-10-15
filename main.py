"""
main.py
=======
Main orchestration pipeline for the 'uttar-pradesh-universities-data-extractor' project.
Coordinates web scraping, data cleaning & validation, rate-limited geocoding,
spatial distance computation (Haversine formula), and export to CSV and styled Excel.

Author: Research Data Analyst Candidate (Web Data and Automation)
"""

import argparse
import logging
import os
import sys
from typing import Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from scraper import scrape_up_universities, DEFAULT_WIKI_URL
from cleaner import clean_university_data, validate_university_data
from geocoder import UniversityGeocoder, GREATER_NOIDA_LAT, GREATER_NOIDA_LON

# Configure root logger
LOG_FORMAT = "%(asctime)s [%(levelname)s] %(name)s: %(message)s"
logging.basicConfig(level=logging.INFO, format=LOG_FORMAT, datefmt="%Y-%m-%d %H:%M:%S")
logger = logging.getLogger("main_pipeline")


def save_as_csv(df: pd.DataFrame, output_path: str) -> None:
    """
    Save DataFrame to CSV with UTF-8 encoding.

    Args:
        df (pd.DataFrame): DataFrame to save.
        output_path (str): File destination path.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    df.to_csv(output_path, index=False, encoding="utf-8")
    logger.info("Successfully exported CSV (%d rows) to: %s", len(df), output_path)


def save_as_styled_excel(df: pd.DataFrame, output_path: str) -> None:
    """
    Save DataFrame as an Excel workbook (.xlsx) with professional styling:
    - Bold header row with a modern dark navy background and white text.
    - Subtle alternating row zebra-striping.
    - Proper cell borders and text alignment.
    - Automatically fitted column widths with generous padding.

    Args:
        df (pd.DataFrame): DataFrame to save.
        output_path (str): File destination path.
    """
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    # 1. Write basic Excel via pandas / openpyxl engine
    df.to_excel(output_path, index=False, engine="openpyxl", sheet_name="UP Universities")

    # 2. Load with openpyxl to apply styling and auto-fit column widths
    wb = openpyxl.load_workbook(output_path)
    ws = wb["UP Universities"]

    # Header styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Thin gray border
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Zebra stripes
    even_row_fill = PatternFill(start_color="F2F4F8", end_color="F2F4F8", fill_type="solid")
    default_font = Font(name="Calibri", size=10)

    # Style Header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Style Data rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=2):
        ws.row_dimensions[row_idx].height = 20
        use_fill = even_row_fill if (row_idx % 2 == 0) else None

        for cell in row:
            cell.font = default_font
            cell.border = thin_border
            if use_fill:
                cell.fill = use_fill

            # Center-align specific numeric and short columns
            header_val = str(ws.cell(row=1, column=cell.column).value or "")
            if header_val in ["established_year", "latitude", "longitude", "distance_from_greater_noida_km", "university_type"]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_length = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_length:
                max_length = len(val_str)
        # Add safety margin
        ws.column_dimensions[col_letter].width = max(max_length + 4, 14)

    # Enable sheet grid lines
    ws.views.sheetView[0].showGridLines = True

    wb.save(output_path)
    logger.info("Successfully exported styled Excel workbook to: %s", output_path)


def run_pipeline(
    wiki_url: str = DEFAULT_WIKI_URL,
    output_dir: str = "output",
    cache_dir: str = "cache",
    delay: float = 1.0,
) -> pd.DataFrame:
    """
    Execute the full end-to-end extraction, transformation, enrichment,
    and export pipeline.

    Args:
        wiki_url (str): Source Wikipedia URL.
        output_dir (str): Destination folder for CSV and Excel files.
        cache_dir (str): Directory storing geocoding cache.
        delay (float): Politeness network delay in seconds.

    Returns:
        pd.DataFrame: Final processed DataFrame.
    """
    logger.info("=================================================================")
    logger.info("STARTING UTTAR PRADESH UNIVERSITIES DATA EXTRACTION PIPELINE")
    logger.info("=================================================================")

    # 1. Scrape raw data from Wikipedia
    logger.info("[STEP 1/4] Scraping raw university tables from Wikipedia...")
    raw_df = scrape_up_universities(url=wiki_url, delay=delay)
    if raw_df.empty:
        logger.error("Scraping returned an empty dataset. Aborting pipeline.")
        sys.exit(1)
    logger.info("Step 1 Complete: Scraped %d raw university records.", len(raw_df))

    # 2. Clean and validate data
    logger.info("[STEP 2/4] Cleaning text, normalizing schema, and validating quality...")
    cleaned_df, metrics = clean_university_data(raw_df)
    is_valid = validate_university_data(cleaned_df, metrics)
    if not is_valid:
        logger.warning("Data quality validation reported warnings/errors. Proceeding with caution.")
    logger.info("Step 2 Complete: Cleaned dataset ready with %d rows.", len(cleaned_df))

    # 3. Geocode and calculate Haversine distance
    logger.info("[STEP 3/4] Geocoding locations and computing distance to Greater Noida...")
    cache_file = os.path.join(cache_dir, "geocode_cache.json")
    geocoder = UniversityGeocoder(cache_path=cache_file)
    enriched_df = geocoder.enrich_dataframe_with_geodata(
        cleaned_df,
        location_col="location",
        ref_lat=GREATER_NOIDA_LAT,
        ref_lon=GREATER_NOIDA_LON,
    )
    logger.info("Step 3 Complete: Geospatial coordinates and distances calculated.")

    # 4. Save to CSV and Excel
    logger.info("[STEP 4/4] Exporting final datasets to output directory...")
    os.makedirs(output_dir, exist_ok=True)
    csv_file = os.path.join(output_dir, "uttar_pradesh_universities.csv")
    excel_file = os.path.join(output_dir, "uttar_pradesh_universities.xlsx")

    save_as_csv(enriched_df, csv_file)
    save_as_styled_excel(enriched_df, excel_file)

    logger.info("=================================================================")
    logger.info("PIPELINE COMPLETED SUCCESSFULLY!")
    logger.info("Total Universities Processed : %d", len(enriched_df))
    logger.info("CSV Export Location         : %s", csv_file)
    logger.info("Excel Export Location       : %s", excel_file)
    logger.info("=================================================================")

    return enriched_df


def main():
    """CLI argument parsing and main execution entry point."""
    parser = argparse.ArgumentParser(
        description="Scrape, clean, geocode, and analyze universities in Uttar Pradesh from Wikipedia."
    )
    parser.add_argument(
        "--url",
        type=str,
        default=DEFAULT_WIKI_URL,
        help="Target Wikipedia URL (defaults to UP universities list).",
    )
    parser.add_argument(
        "--output-dir",
        type=str,
        default="output",
        help="Folder where CSV and Excel output files will be saved.",
    )
    parser.add_argument(
        "--cache-dir",
        type=str,
        default="cache",
        help="Folder where geocoding cache JSON will be stored.",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=1.0,
        help="Polite request delay in seconds (default: 1.0).",
    )

    args = parser.parse_args()

    # Determine paths relative to this script file
    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(base_dir, args.output_dir)
    cache_dir = os.path.join(base_dir, args.cache_dir)

    run_pipeline(
        wiki_url=args.url,
        output_dir=output_dir,
        cache_dir=cache_dir,
        delay=args.delay,
    )


if __name__ == "__main__":
    main()
